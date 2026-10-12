# ============================================================
# Test des 10 exigences sélectionnées par le manager
# ============================================================

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
from time import perf_counter
from typing import Any, Callable

import pandas as pd

from src.rule_loader import load_rules
from src.rule_engine import analyse_requirements
from src.llm_analyzer import LLMAnalyzer, LLMSettings


MANAGER_REQUIREMENT_IDS = [
    "PKP-ANNEX1.CUS-72",
    "PKP-ANNEX1.CUS-198",
    "PKP-ANNEX1.CUS-3475",
    "PKP-ANNEX1.CUS-235",
    "PKP-ANNEX1.CUS-237",
    "PKP-ANNEX1.CUS-3482",
    "PKP-ANNEX1.CUS-3489",
    "PKP-ANNEX1.CUS-314",
    "PKP-ANNEX1.CUS-562",
    "PKP-ANNEX1.CUS-609",
]

DEFAULT_RULES_PATH = Path("rules/rules_en.csv")
DEFAULT_OUTPUT_DIRECTORY = Path("data/output")
DEFAULT_MODEL = "qwen2.5:3b"


# ============================================================
# Chargement Excel compatible avec les versions précédentes
# ============================================================

def get_excel_loader() -> Callable[[str], dict[str, Any]]:
    """Retourne la fonction existante de chargement Excel."""

    try:
        from src.excel_loader import load_requirements
        return load_requirements
    except (ImportError, ModuleNotFoundError):
        pass

    try:
        from src.excel_loader import load_required_columns
        return load_required_columns
    except (ImportError, ModuleNotFoundError):
        pass

    try:
        from load_excel import load_required_columns
        return load_required_columns
    except (ImportError, ModuleNotFoundError) as error:
        raise ImportError(
            "Aucun chargeur Excel compatible n'a été trouvé."
        ) from error


# ============================================================
# Arguments
# ============================================================

def parse_arguments() -> argparse.Namespace:
    """Lit les arguments de la ligne de commande."""

    parser = argparse.ArgumentParser(
        description=(
            "Teste le moteur de règles et Qwen2.5:3b sur les "
            "10 exigences sélectionnées par le manager."
        )
    )

    parser.add_argument(
        "--input",
        "-i",
        required=True,
        help="Chemin de l'export Excel DOORS.",
    )

    parser.add_argument(
        "--rules",
        "-r",
        default=str(DEFAULT_RULES_PATH),
        help="Chemin du fichier rules_en.csv.",
    )

    parser.add_argument(
        "--output",
        "-o",
        help="Chemin du rapport Excel de comparaison.",
    )

    parser.add_argument(
        "--model",
        default=DEFAULT_MODEL,
        help="Modèle Ollama à utiliser.",
    )

    parser.add_argument(
        "--skip-llm",
        action="store_true",
        help="Exécute uniquement le moteur de règles.",
    )

    return parser.parse_args()


def clean_path(value: str) -> Path:
    """Nettoie un chemin reçu depuis PowerShell."""

    return Path(
        value.strip().strip('"').strip("'")
    ).expanduser().resolve()


# ============================================================
# Préparation des constats pour le LLM
# ============================================================

def findings_for_requirement(
    requirement_id: str,
    findings_dataframe: pd.DataFrame,
) -> list[dict[str, str]]:
    """Retourne les anomalies déterministes d'une exigence."""

    if findings_dataframe.empty:
        return []

    selected = findings_dataframe[
        findings_dataframe["requirement_id"].astype(str)
        == str(requirement_id)
    ]

    findings: list[dict[str, str]] = []

    for _, row in selected.iterrows():
        findings.append(
            {
                "rule_id": str(row.get("rule_id", "")).strip(),
                "category": str(row.get("category", "")).strip(),
                "severity": str(row.get("severity", "")).strip(),
                "evidence": str(row.get("evidence", "")).strip(),
                "explanation": str(
                    row.get("explanation", "")
                ).strip(),
            }
        )

    return findings


def build_full_review_finding() -> list[dict[str, str]]:
    """
    Crée un signal de revue générale lorsque les règles lexicales
    ne détectent rien.

    Ce signal ne prétend pas qu'une anomalie existe. Il demande au
    LLM de déterminer si l'exigence présente un problème contextuel.
    """

    return [
        {
            "rule_id": "FULL-REVIEW-001",
            "category": "FULL_CONTEXT_REVIEW",
            "severity": "MEDIUM",
            "evidence": "complete requirement statement",
            "explanation": (
                "Perform a full contextual quality review. Set "
                "issue_present to true only if the requirement is "
                "ambiguous, incomplete, non-verifiable, non-atomic, "
                "weakly binding, or dependent on undefined conditions."
            ),
        }
    ]


# ============================================================
# Affichage de progression
# ============================================================

def print_progress(
    current: int,
    total: int,
    requirement_id: str,
    elapsed_seconds: float,
) -> None:
    """Affiche la progression et une estimation restante."""

    average_seconds = elapsed_seconds / current
    remaining_seconds = average_seconds * (total - current)

    print(
        f"[{current}/{total}] {requirement_id} - "
        f"écoulé: {elapsed_seconds:.1f}s - "
        f"reste estimé: {remaining_seconds:.1f}s"
    )


# ============================================================
# Export Excel
# ============================================================

def format_worksheet(worksheet) -> None:
    """Applique une mise en forme simple au rapport."""

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(color="FFFFFF", bold=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    for column_number, cells in enumerate(
        worksheet.iter_cols(),
        start=1,
    ):
        maximum_length = 0

        for cell in cells[:30]:
            value = "" if cell.value is None else str(cell.value)
            maximum_length = max(
                maximum_length,
                max(
                    (len(line) for line in value.splitlines()),
                    default=0,
                ),
            )

        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = min(max(maximum_length + 2, 12), 60)


def export_report(
    output_path: Path,
    comparison_dataframe: pd.DataFrame,
    findings_dataframe: pd.DataFrame,
    run_information: pd.DataFrame,
) -> None:
    """Exporte la comparaison manager, règles et LLM."""

    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:
        comparison_dataframe.to_excel(
            writer,
            sheet_name="Manager_Test_Results",
            index=False,
        )

        findings_dataframe.to_excel(
            writer,
            sheet_name="Rule_Findings",
            index=False,
        )

        run_information.to_excel(
            writer,
            sheet_name="Run_Information",
            index=False,
        )

        for worksheet in writer.book.worksheets:
            format_worksheet(worksheet)


# ============================================================
# Exécution du test
# ============================================================

def run_manager_test(
    input_path: Path,
    rules_path: Path,
    output_path: Path,
    model: str,
    use_llm: bool,
) -> None:
    """Exécute le test complet sur les dix exigences."""

    global_started_at = perf_counter()

    excel_loader = get_excel_loader()
    excel_result = excel_loader(str(input_path))
    requirements = excel_result["requirements"].copy()

    manager_requirements = requirements[
        requirements["ID"].astype(str).isin(
            MANAGER_REQUIREMENT_IDS
        )
    ].copy()

    manager_requirements["_manager_order"] = (
        manager_requirements["ID"]
        .astype(str)
        .map(
            {
                requirement_id: position
                for position, requirement_id
                in enumerate(MANAGER_REQUIREMENT_IDS)
            }
        )
    )

    manager_requirements = manager_requirements.sort_values(
        "_manager_order"
    ).drop(columns=["_manager_order"])

    found_ids = set(manager_requirements["ID"].astype(str))
    missing_ids = [
        requirement_id
        for requirement_id in MANAGER_REQUIREMENT_IDS
        if requirement_id not in found_ids
    ]

    if missing_ids:
        raise ValueError(
            "Les ID suivants sont absents de l'export : "
            + ", ".join(missing_ids)
        )

    rules_result = load_rules(
        rules_path=rules_path,
        include_disabled=False,
        strict=True,
    )

    rule_result = analyse_requirements(
        requirements=manager_requirements,
        rules=rules_result["rules"],
    )

    summary = rule_result["summary"].copy()
    findings_dataframe = rule_result["findings"].copy()

    summary["Manager_Expected_Flag"] = "YES"
    summary["Rule_Engine_Flag"] = summary["Issue_Count"].apply(
        lambda count: "YES" if int(count) > 0 else "NO"
    )
    summary["Review_Mode"] = ""
    summary["LLM_Analyzed"] = "NO"
    summary["LLM_Model"] = model if use_llm else ""
    summary["LLM_Decision"] = "NOT_RUN"
    summary["LLM_Confirmed_Findings"] = ""
    summary["LLM_Rejected_Findings"] = ""
    summary["LLM_Missing_Information"] = ""
    summary["LLM_Clarification_Questions"] = ""
    summary["LLM_Proposed_Rewrite"] = ""
    summary["LLM_Risk_Recommended"] = ""
    summary["LLM_Risk_Title"] = ""
    summary["LLM_Risk_Cause"] = ""
    summary["LLM_Risk_Event"] = ""
    summary["LLM_Risk_Consequences"] = ""
    summary["LLM_Processing_Time_s"] = 0.0
    summary["LLM_Prompt_Tokens"] = 0
    summary["LLM_Generated_Tokens"] = 0
    summary["LLM_Tokens_Per_Second"] = 0.0
    summary["LLM_Error"] = ""

    if use_llm:
        settings = LLMSettings(
            model=model,
            temperature=0.1,
            num_ctx=4096,
            num_predict=500,
            seed=42,
            keep_alive="30m",
        )
        analyzer = LLMAnalyzer(settings=settings)
        health = analyzer.health_check()

        if not health["model_available"]:
            raise ValueError(
                f"Le modèle {model} n'est pas installé. "
                f"Exécute : ollama pull {model}"
            )

        llm_batch_started_at = perf_counter()
        total = len(summary)

        for position, row_index in enumerate(
            summary.index,
            start=1,
        ):
            requirement_id = str(summary.at[row_index, "ID"])
            original_text = str(
                summary.at[
                    row_index,
                    "Object Heading and Object Text",
                ]
            )
            object_type = str(
                summary.at[row_index, "ObjectType"]
            )

            requirement_findings = findings_for_requirement(
                requirement_id=requirement_id,
                findings_dataframe=findings_dataframe,
            )

            if requirement_findings:
                review_mode = "VALIDATE_RULE_FINDINGS"
            else:
                review_mode = "FULL_CONTEXT_REVIEW"
                requirement_findings = build_full_review_finding()

            analysis = analyzer.analyze_requirement(
                requirement_id=requirement_id,
                original_text=original_text,
                object_type=object_type,
                findings=requirement_findings,
            )

            summary.at[row_index, "Review_Mode"] = review_mode
            summary.at[row_index, "LLM_Analyzed"] = (
                "YES" if analysis.analyzed else "NO"
            )
            summary.at[row_index, "LLM_Decision"] = analysis.decision
            summary.at[
                row_index,
                "LLM_Confirmed_Findings",
            ] = str(analysis.confirmed_findings)
            summary.at[
                row_index,
                "LLM_Rejected_Findings",
            ] = str(analysis.rejected_findings)
            summary.at[
                row_index,
                "LLM_Missing_Information",
            ] = " | ".join(analysis.missing_information)
            summary.at[
                row_index,
                "LLM_Clarification_Questions",
            ] = " | ".join(analysis.clarification_questions)
            summary.at[
                row_index,
                "LLM_Proposed_Rewrite",
            ] = analysis.proposed_rewrite
            summary.at[
                row_index,
                "LLM_Risk_Recommended",
            ] = "YES" if analysis.risk_recommended else "NO"
            summary.at[
                row_index,
                "LLM_Risk_Title",
            ] = analysis.risk_draft.title
            summary.at[
                row_index,
                "LLM_Risk_Cause",
            ] = analysis.risk_draft.cause
            summary.at[
                row_index,
                "LLM_Risk_Event",
            ] = analysis.risk_draft.event
            summary.at[
                row_index,
                "LLM_Risk_Consequences",
            ] = " | ".join(analysis.risk_draft.consequences)
            summary.at[
                row_index,
                "LLM_Processing_Time_s",
            ] = analysis.processing_time_seconds
            summary.at[
                row_index,
                "LLM_Prompt_Tokens",
            ] = analysis.prompt_tokens
            summary.at[
                row_index,
                "LLM_Generated_Tokens",
            ] = analysis.generated_tokens
            summary.at[
                row_index,
                "LLM_Tokens_Per_Second",
            ] = analysis.tokens_per_second
            summary.at[
                row_index,
                "LLM_Error",
            ] = analysis.error

            print_progress(
                current=position,
                total=total,
                requirement_id=requirement_id,
                elapsed_seconds=(
                    perf_counter() - llm_batch_started_at
                ),
            )

    elapsed_seconds = perf_counter() - global_started_at

    positive_rule_count = int(
        (summary["Rule_Engine_Flag"] == "YES").sum()
    )
    rule_recall = positive_rule_count / len(MANAGER_REQUIREMENT_IDS)

    llm_positive_decisions = {
        "CONFIRMED",
        "PARTIALLY_CONFIRMED",
        "NEEDS_EXPERT_REVIEW",
    }
    llm_detected_count = int(
        summary["LLM_Decision"].isin(
            llm_positive_decisions
        ).sum()
    ) if use_llm else 0
    llm_recall = (
        llm_detected_count / len(MANAGER_REQUIREMENT_IDS)
        if use_llm
        else 0.0
    )

    run_information = pd.DataFrame(
        [
            {"Metric": "Input file", "Value": str(input_path)},
            {"Metric": "Rules file", "Value": str(rules_path)},
            {"Metric": "Model", "Value": model if use_llm else "Not used"},
            {"Metric": "Manager cases", "Value": len(MANAGER_REQUIREMENT_IDS)},
            {"Metric": "Cases found", "Value": len(summary)},
            {"Metric": "Rule engine detected", "Value": positive_rule_count},
            {"Metric": "Rule engine recall", "Value": rule_recall},
            {"Metric": "LLM detected or review", "Value": llm_detected_count},
            {"Metric": "LLM recall", "Value": llm_recall},
            {"Metric": "Execution time seconds", "Value": round(elapsed_seconds, 2)},
            {
                "Metric": "Average LLM time seconds",
                "Value": round(
                    summary["LLM_Processing_Time_s"].mean(),
                    2,
                ) if use_llm else 0.0,
            },
            {
                "Metric": "Execution date",
                "Value": datetime.now().isoformat(timespec="seconds"),
            },
        ]
    )

    export_report(
        output_path=output_path,
        comparison_dataframe=summary,
        findings_dataframe=findings_dataframe,
        run_information=run_information,
    )

    print("\n" + "=" * 65)
    print("RÉSULTAT DU TEST MANAGER")
    print("=" * 65)
    print(f"Exigences trouvées       : {len(summary)}/10")
    print(f"Détectées par les règles : {positive_rule_count}/10")
    print(f"Rappel moteur de règles  : {rule_recall:.1%}")

    if use_llm:
        print(f"Signalées par le LLM     : {llm_detected_count}/10")
        print(f"Rappel LLM               : {llm_recall:.1%}")

    print(f"Temps total              : {elapsed_seconds:.2f} s")
    print(f"Rapport                  : {output_path}")


# ============================================================
# Main
# ============================================================

def main() -> None:
    """Point d'entrée du script."""

    arguments = parse_arguments()

    input_path = clean_path(arguments.input)
    rules_path = clean_path(arguments.rules)

    if arguments.output:
        output_path = clean_path(arguments.output)
    else:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = (
            DEFAULT_OUTPUT_DIRECTORY
            / f"manager_10_requirements_test_{timestamp}.xlsx"
        ).resolve()

    if not input_path.exists():
        raise FileNotFoundError(
            f"Fichier d'entrée introuvable : {input_path}"
        )

    if not rules_path.exists():
        raise FileNotFoundError(
            f"Catalogue de règles introuvable : {rules_path}"
        )

    run_manager_test(
        input_path=input_path,
        rules_path=rules_path,
        output_path=output_path,
        model=arguments.model,
        use_llm=not arguments.skip_llm,
    )


if __name__ == "__main__":
    main()
